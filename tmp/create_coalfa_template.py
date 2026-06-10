import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Solicitud Insumos"

# Estilos básicos
bold_font = Font(bold=True)
header_font = Font(color="FFFFFF", bold=True, size=10)
title_font = Font(bold=True, size=14, color="003366")
blue_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
light_blue_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_alignment = Alignment(horizontal="left", vertical="center")

thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

# Diseño Cabecera
ws.merge_cells('A1:B2')
ws['A1'] = "COALFA"
ws['A1'].font = title_font
ws['A1'].alignment = center_alignment

ws.merge_cells('C1:F2')
ws['C1'] = "Registro de solicitud de insumos"
ws['C1'].font = Font(bold=True, size=12)
ws['C1'].alignment = center_alignment

# Fila 3: Solicitado por
ws.merge_cells('A3:B3')
ws['A3'] = "Solicitado por:"
ws['A3'].font = bold_font
ws['A3'].fill = light_blue_fill
ws['A3'].border = thin_border

ws.merge_cells('C3:F3')
ws['C3'] = "REPOSTERIA" # Area Placeholder
ws['C3'].font = bold_font
ws['C3'].alignment = Alignment(horizontal="left", vertical="center")
ws['C3'].border = thin_border

# Fila 4: Solicitante / Servicio
ws['A4'] = "Solicitante:"
ws['A4'].font = bold_font
ws['A4'].border = thin_border
ws['B4'] = ""
ws['B4'].border = thin_border
ws.merge_cells('C4:D4')
ws['C4'] = ""  # Name
ws['C4'].border = thin_border
ws['E4'] = "Servicio:"
ws['E4'].font = bold_font
ws['E4'].border = thin_border
ws['F4'] = ""
ws['F4'].border = thin_border

# Fila 5: Fecha solicitud / Fecha de servicio
ws['A5'] = "Fecha solicitud:"
ws['A5'].font = bold_font
ws['A5'].border = thin_border
ws.merge_cells('B5:D5')
ws['B5'] = ""
ws['B5'].border = thin_border
ws['E5'] = "Fecha servicio:"
ws['E5'].font = bold_font
ws['E5'].border = thin_border
ws['F5'] = ""
ws['F5'].border = thin_border

# Fila 6: Raciones
ws.merge_cells('A6:C6')
ws['A6'] = "Cantidad proyectada por comensales:"
ws['A6'].font = bold_font
ws['A6'].border = thin_border
ws.merge_cells('D6:F6')
ws['D6'] = ""
ws['D6'].border = thin_border

# Separador
ws.row_dimensions[7].height = 5

# Fila 8: Cabeceras de la tabla
headers = [
    "INSUMOS / PLANTA",
    "Cantidad solicitada", 
    "Cantidad entregada",
    "F. V.",
    "Lote",
    "M / HF"
]

for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=8, column=col_idx)
    cell.value = header
    cell.font = header_font
    cell.fill = blue_fill
    cell.alignment = center_alignment
    cell.border = thin_border

# Anchos de columna
ws.column_dimensions['A'].width = 30 # Insumos
ws.column_dimensions['B'].width = 18 # Cantidad Solicitada
ws.column_dimensions['C'].width = 18 # Cantidad Entregada
ws.column_dimensions['D'].width = 15 # F.V
ws.column_dimensions['E'].width = 20 # Lote
ws.column_dimensions['F'].width = 15 # M / HF

# Añadir filas de estilo en blanco
for r in range(9, 25):
    for c in range(1, 7):
        ws.cell(row=r, column=c).border = thin_border
    # Alinear texto de insumos
    ws.cell(row=r, column=1).alignment = left_alignment
    ws.cell(row=r, column=3).alignment = center_alignment

# Ejemplo visual
ws['A9'] = "durazno"
ws['B9'] = "2 K"
ws['C9'] = 3
ws['A10'] = "Huevos"
ws['B10'] = "30 u"
ws['C10'] = 30

file_path = r"c:\Users\Coalfa\Desktop\Formato_Solicitud_COALFA.xlsx"
wb.save(file_path)
print(f"Archivo guardado exitosamente en: {file_path}")
