import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import os

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Solicitudes"

# Encabezados
headers = [
    "CONCEPTO", 
    "ARTÍCULO", 
    "SKU", 
    "UM", 
    "SALDO FISICO STOCK", 
    "COSTO UNITARIO", 
    "COSTO TOTAL", 
    "SALIDA DIARIA DE BODEGA", 
    "STOCK EN BODEGA", 
    "OBSERVACION"
]

# Estilos
header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
center_alignment = Alignment(horizontal="center", vertical="center")

ws.append(headers)

# Aplicar estilos a cabecera
for col in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center_alignment

# Anchos de columna para que se vea bien
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 30 # Articulo
ws.column_dimensions['C'].width = 10
ws.column_dimensions['D'].width = 10
ws.column_dimensions['E'].width = 20
ws.column_dimensions['F'].width = 15
ws.column_dimensions['G'].width = 15
ws.column_dimensions['H'].width = 25 # Salida Diaria
ws.column_dimensions['I'].width = 18
ws.column_dimensions['J'].width = 25 # Observacion

# Añadir filas de ejemplo (la lógica lo ignora si Salida Diaria es vacía o cero)
ws.append(["Ejemplo", "ACEITE DE 5 LT.", "ABA-001", "UN", "", "", "", 2, "", "Cocina Caliente"])
ws.append(["Demostracion", "Sal de Mar", "ABA-002", "KG", "", "", "", 1.5, "", "Cocina Fría"])

# Guardar en local
file_path = r"c:\Users\Coalfa\Desktop\Formato_Subida_Solicitudes.xlsx"
wb.save(file_path)
print(f"Archivo creado en: {file_path}")
